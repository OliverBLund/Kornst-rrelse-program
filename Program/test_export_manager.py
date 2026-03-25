"""
Regression tests for export manager CSV filtering behavior.
"""

import csv
import json
import os
import sys
import tempfile
import unittest

sys.path.insert(0, 'Program')

from data_loader import GrainSizeData
from gui.export_manager import ExportManager
from k_calculations_v2 import CalculationStatus, KCalculationResult


def build_dataset(name: str = 'Sample A') -> GrainSizeData:
    """Create a stable grain-size dataset for export tests."""
    return GrainSizeData(
        sample_name=name,
        temperature=20.0,
        porosity=0.35,
        particle_sizes=[4.75, 2.0, 1.0, 0.5, 0.25, 0.125, 0.063],
        percent_passing=[100.0, 95.0, 82.0, 61.0, 38.0, 14.0, 4.0],
    )


def build_results() -> list[KCalculationResult]:
    """Create representative K-results spanning multiple method groups."""
    return [
        KCalculationResult(
            method_name='Hazen',
            k_value=1.0e-4,
            formula_used='k = c * d10^2',
            status=CalculationStatus.OK,
            status_message='',
            conditions_met=True,
            temperature=20.0,
            porosity=0.35,
            grain_size_used='D10',
        ),
        KCalculationResult(
            method_name='Beyer',
            k_value=1.5e-4,
            formula_used='k = c * d10^2 * log(500 / Cu)',
            status=CalculationStatus.OK,
            status_message='',
            conditions_met=True,
            temperature=20.0,
            porosity=0.35,
            grain_size_used='D10',
        ),
        KCalculationResult(
            method_name='USBR',
            k_value=2.5e-4,
            formula_used='k = c * d20^2.3',
            status=CalculationStatus.OK,
            status_message='',
            conditions_met=True,
            temperature=20.0,
            porosity=0.35,
            grain_size_used='D20',
        ),
    ]


def read_csv_rows(path: str) -> list[list[str]]:
    with open(path, newline='', encoding='utf-8') as handle:
        return list(csv.reader(handle))


def read_json(path: str) -> dict:
    with open(path, encoding='utf-8') as handle:
        return json.load(handle)


class TestExportManagerExports(unittest.TestCase):
    def setUp(self):
        self.dataset = build_dataset()
        self.results = build_results()
        self.datasets = [(self.dataset.sample_name, self.dataset, self.results)]

    def make_config(self, output_dir: str, **overrides):
        config = {
            'csv': True,
            'csv_mode': 'combined',
            'csv_long': True,
            'csv_wide': False,
            'excel': False,
            'json': False,
            'png': False,
            'svg': False,
            'pdf_plot': False,
            'grain_distribution': False,
            'percentiles': True,
            'gradation': True,
            'classification': False,
            'k_values': True,
            'statistics': False,
            'plots': False,
            'formulas': False,
            'validation': False,
            'k_filter_mode': 'all',
            'selected_k_categories': {
                'hazen_based': True,
                'porosity_dependent': True,
                'uniformity_dependent': True,
                'empirical': True,
                'temperature_corrected': True,
            },
            'selected_percentiles': ['d10', 'd20', 'd30', 'd50', 'd60'],
            'selected_k_methods': None,
            'k_units': {
                'm_s': True,
                'cm_s': True,
                'm_d': True,
            },
            'selected_statistics': ['mean', 'median', 'std_dev', 'min', 'max', 'valid_count'],
            'include_grain_size_stats': True,
            'include_metadata': {
                'sample_info': True,
                'environmental': True,
                'export_timestamp': True,
            },
            'output_dir': output_dir,
            'filename_template': '{sample_name}_results_{date}',
        }
        config.update(overrides)
        return config

    def test_combined_long_csv_respects_selected_methods_and_units(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            config = self.make_config(
                temp_dir,
                csv_long=True,
                csv_wide=False,
                k_filter_mode='individual',
                selected_k_methods=['Beyer', 'USBR'],
                k_units={'m_s': True, 'cm_s': False, 'm_d': True},
            )

            exported = ExportManager().export(self.datasets, config)

            self.assertEqual(len(exported), 1)
            export_name = os.path.basename(exported[0])
            self.assertIn('combined_all_datasets', export_name)

            rows = read_csv_rows(exported[0])
            header = rows[0]
            methods = [row[1] for row in rows[1:]]

            self.assertIn('K (m/s)', header)
            self.assertIn('K (m/d)', header)
            self.assertNotIn('K (cm/s)', header)
            self.assertEqual(methods, ['Beyer', 'USBR'])
            self.assertTrue(all(len(row) == len(header) for row in rows[1:]))

    def test_wide_csv_only_writes_requested_format_and_category_methods(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            config = self.make_config(
                temp_dir,
                csv_long=False,
                csv_wide=True,
                k_filter_mode='category',
                selected_k_categories={
                    'hazen_based': False,
                    'porosity_dependent': False,
                    'uniformity_dependent': True,
                    'empirical': False,
                    'temperature_corrected': False,
                },
                selected_percentiles=['d10', 'd60'],
                k_units={'m_s': False, 'cm_s': True, 'm_d': False},
                selected_statistics=['median', 'valid_count'],
            )

            exported = ExportManager().export(self.datasets, config)

            self.assertEqual(len(exported), 1)
            export_name = os.path.basename(exported[0])
            self.assertIn('wide_format_all_datasets', export_name)

            rows = read_csv_rows(exported[0])
            header = rows[0]
            data_row = rows[1]

            self.assertIn('D10_mm', header)
            self.assertIn('D60_mm', header)
            self.assertNotIn('D20_mm', header)
            self.assertIn('K_Beyer_cm/s', header)
            self.assertIn('Status_Beyer', header)
            self.assertIn('K_Median_cm/s', header)
            self.assertNotIn('K_Mean_cm/s', header)
            self.assertNotIn('K_Hazen_cm/s', header)
            self.assertNotIn('K_Beyer_m/s', header)
            self.assertNotIn('K_Beyer_m/d', header)
            self.assertEqual(data_row[header.index('Valid_Methods_Count')], '1')

    def test_separate_statistics_csv_uses_filtered_results_metadata_and_stat_selection(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            config = self.make_config(
                temp_dir,
                csv_mode='separate',
                csv_long=True,
                csv_wide=False,
                grain_distribution=False,
                k_values=False,
                statistics=True,
                k_filter_mode='individual',
                selected_k_methods=['USBR'],
                k_units={'m_s': False, 'cm_s': False, 'm_d': True},
                selected_statistics=['median', 'valid_count'],
                include_grain_size_stats=False,
                include_metadata={
                    'sample_info': False,
                    'environmental': False,
                    'export_timestamp': True,
                },
            )

            exported = ExportManager().export(self.datasets, config)

            self.assertEqual(len(exported), 1)
            export_name = os.path.basename(exported[0])
            self.assertIn('statistics', export_name)

            rows = read_csv_rows(exported[0])
            timestamp_rows = [row for row in rows if row and row[0] == 'Export Timestamp']
            self.assertEqual(len(timestamp_rows), 1)
            self.assertTrue(timestamp_rows[0][1])
            self.assertIn(['Statistic', 'K (m/d)'], rows)
            self.assertIn(['Median', '21.60'], rows)
            self.assertIn(['Valid Count', '1'], rows)
            self.assertNotIn(['Mean', '21.60'], rows)
            self.assertNotIn(['Sample Name', self.dataset.sample_name], rows)
            self.assertFalse(any(row and row[0] == 'Grain Size Percentiles' for row in rows))

    def test_json_export_honors_metadata_units_and_selected_statistics(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            config = self.make_config(
                temp_dir,
                csv=False,
                json=True,
                statistics=True,
                k_filter_mode='individual',
                selected_k_methods=['USBR'],
                k_units={'m_s': False, 'cm_s': True, 'm_d': False},
                selected_statistics=['median', 'valid_count'],
                include_metadata={
                    'sample_info': False,
                    'environmental': False,
                    'export_timestamp': True,
                },
            )

            exported = ExportManager().export(self.datasets, config)

            self.assertEqual(len(exported), 1)
            payload = read_json(exported[0])

            self.assertNotIn('sample_name', payload)
            self.assertEqual(set(payload['metadata'].keys()), {'exported_at'})
            self.assertTrue(payload['metadata']['exported_at'])
            self.assertEqual(len(payload['k_values']), 1)
            self.assertEqual(payload['k_values'][0]['method'], 'USBR')
            self.assertIn('k_cm_s', payload['k_values'][0])
            self.assertNotIn('k_m_s', payload['k_values'][0])
            self.assertNotIn('k_m_d', payload['k_values'][0])
            self.assertEqual(set(payload['statistics'].keys()), {'median_k_cm_s', 'valid_count'})

    def test_excel_export_honors_metadata_and_stat_selection(self):
        try:
            from openpyxl import load_workbook
        except ImportError:
            self.skipTest('openpyxl not installed')

        with tempfile.TemporaryDirectory() as temp_dir:
            config = self.make_config(
                temp_dir,
                csv=False,
                excel=True,
                statistics=True,
                k_units={'m_s': False, 'cm_s': False, 'm_d': True},
                selected_statistics=['median', 'valid_count'],
                include_grain_size_stats=False,
                include_metadata={
                    'sample_info': False,
                    'environmental': True,
                    'export_timestamp': True,
                },
            )

            exported = ExportManager().export(self.datasets, config)

            self.assertEqual(len(exported), 1)
            workbook = load_workbook(exported[0], data_only=True)
            summary_values = [row[0] for row in workbook['Summary'].iter_rows(values_only=True) if row and row[0]]
            stats_rows = list(workbook['Statistics'].iter_rows(values_only=True))

            self.assertNotIn('Sample Name:', summary_values)
            self.assertIn('Temperature:', summary_values)
            self.assertIn('Exported At:', summary_values)
            self.assertIn('Median (m/d):', summary_values)
            self.assertNotIn('Mean (m/d):', summary_values)
            self.assertEqual(stats_rows[0], ('Statistic', 'K (m/d)'))
            self.assertIn(('Median', 12.96), stats_rows)
            self.assertIn(('Valid Count', 3), stats_rows)
            self.assertFalse(any(row and row[0] == 'Mean' for row in stats_rows[1:]))


if __name__ == '__main__':
    unittest.main(verbosity=2)
