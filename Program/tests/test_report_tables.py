"""Regression tests for report table layout and companion Excel output."""

import io
import sys
import zipfile

sys.path.insert(0, "Program")

from openpyxl import load_workbook

from data_loader import GrainSizeData
from k_calculations import CalculationStatus, KCalculationResult
from report_generator import ReportGenerator
from report_tables import analyze_report_tables, externalize_report_tables, extract_report_tables, generate_excel_appendix


def build_dataset(name: str) -> GrainSizeData:
    return GrainSizeData(
        sample_name=name,
        temperature=20.0,
        porosity=0.35,
        particle_sizes=[4.75, 2.0, 1.0, 0.5, 0.25, 0.125, 0.063],
        percent_passing=[100.0, 95.0, 82.0, 61.0, 38.0, 14.0, 4.0],
    )


def build_results(multiplier: float = 1.0) -> list[KCalculationResult]:
    return [
        KCalculationResult(
            method_name=method,
            k_value=(index + 1) * 1.0e-5 * multiplier,
            formula_used=f"formula {method}",
            status=CalculationStatus.OK,
            status_message="",
            conditions_met=True,
            temperature=20.0,
            porosity=0.35,
            grain_size_used="D10",
        )
        for index, method in enumerate(("Hazen", "Beyer", "USBR"))
    ]


def comparison_html(sample_count: int) -> str:
    datasets = [build_dataset(f"Sample {index:02d}") for index in range(1, sample_count + 1)]
    details = [
        {
            "label": dataset.sample_name,
            "dataset": dataset,
            "k_results": build_results(1.0 + index / 10),
            "temperature": dataset.temperature,
            "porosity": dataset.porosity,
            "group_name": "Ungrouped",
        }
        for index, dataset in enumerate(datasets)
    ]
    return ReportGenerator().generate_comparison_report(
        datasets,
        sample_details=details,
        sections={
            "cover_page": False,
            "executive_summary": False,
            "methodology": False,
            "results": True,
            "plots": False,
            "interpretation": False,
            "grain_comparison": True,
            "k_statistics": True,
        },
    )


def test_large_comparison_tables_are_portrait_friendly_and_recommend_excel():
    analysis = analyze_report_tables(comparison_html(20))
    by_id = {table.table_id: table for table in analysis.tables}

    assert by_id["sample-properties"].column_count == 7
    assert by_id["sample-classification"].column_count == 4  # +Descriptor column
    assert by_id["k-aggregate-statistics"].column_count == 6
    assert by_id["k-aggregate-coverage"].column_count == 4
    assert by_id["k-method-results-comparison"].column_count == 6
    assert max(table.column_count for table in analysis.tables) <= 7
    assert analysis.excel_recommended
    assert {
        "grain-parameters",
        "k-method-results-comparison",
    }.issubset({table.table_id for table in analysis.large_tables})


def test_small_comparison_does_not_recommend_excel():
    analysis = analyze_report_tables(comparison_html(5))
    assert not analysis.excel_recommended


def test_externalize_report_tables_replaces_only_selected_tables():
    html = """
    <html><body>
    <h3>Small Table</h3>
    <table data-report-table="small-table"><tr><td>keep me</td></tr></table>
    <h3>Large Table</h3>
    <table data-report-table="large-table"><tr><td>remove<br>me</td></tr></table>
    <p>after table</p>
    </body></html>
    """

    rewritten = externalize_report_tables(html, {"large-table": "Large Table"})

    assert "keep me" in rewritten
    assert "remove" not in rewritten
    assert "after table" in rewritten
    assert "Large table moved to companion Excel appendix" in rewritten
    assert 'data-externalized-table="large-table"' in rewritten


def test_excel_appendix_preserves_colspan_and_uses_table_titles():
    html = """
    <h3>Method Summary</h3>
    <table data-report-table="method-summary">
      <thead><tr><th>Method</th><th>Value</th><th>Status</th></tr></thead>
      <tbody>
        <tr><td colspan="2">No results</td><td>Unavailable</td></tr>
      </tbody>
    </table>
    """
    workbook = load_workbook(io.BytesIO(generate_excel_appendix(html)))
    sheet = workbook["Method Summary"]

    assert "A2:B2" in {str(cell_range) for cell_range in sheet.merged_cells.ranges}
    assert sheet["A2"].value == "No results"
    assert sheet.sheet_view.showGridLines is False


def test_docx_renderer_preserves_colspan():
    html = """
    <html><body>
    <h3>Method Summary</h3>
    <table data-report-table="method-summary">
      <thead><tr><th>Method</th><th>Value</th><th>Status</th></tr></thead>
      <tbody>
        <tr><td colspan="2">No results</td><td>Unavailable</td></tr>
      </tbody>
    </table>
    </body></html>
    """
    blob = ReportGenerator().generate_docx_from_html(html)
    with zipfile.ZipFile(io.BytesIO(blob)) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")

    assert "w:gridSpan" in document_xml
    assert "No results" in document_xml


def test_docx_renderer_adds_report_chrome_and_box_styling():
    html = """
    <html><body>
    <h1>Grain Size Analysis Report</h1>
    <div class="success-box"><p>Executive summary text</p></div>
    <div class="info-box"><h3>Methodology</h3><p>Method text</p></div>
    </body></html>
    """

    blob = ReportGenerator().generate_docx_from_html(
        html,
        metadata={
            "project_name": "Project Alpha",
            "project_no": "P-42",
            "client": "Client Beta",
            "analyst": "Analyst Gamma",
            "date": "2026-06-30",
        },
    )
    with zipfile.ZipFile(io.BytesIO(blob)) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
        header_xml = archive.read("word/header1.xml").decode("utf-8")
        footer_xml = archive.read("word/footer1.xml").decode("utf-8")

    assert "Executive summary text" in document_xml
    assert "Methodology" in document_xml
    assert 'w:fill="F4F9F4"' in document_xml
    assert 'w:fill="F7F7F7"' in document_xml
    assert "Project Alpha" in header_xml
    assert "P-42" in header_xml
    assert "Analyst Gamma" in footer_xml
    assert "Client Beta" in footer_xml

def test_docx_renderer_replaces_externalized_tables_with_appendix_note():
    rows = "".join(
        f"<tr><td>Sample {index}</td><td>{index}</td></tr>"
        for index in range(60)
    )
    html = f"""
    <html><body>
    <h3>Large Result Table</h3>
    <table data-report-table="large-result-table">
      <thead><tr><th>Sample</th><th>Value</th></tr></thead>
      <tbody>{rows}</tbody>
    </table>
    </body></html>
    """

    blob = ReportGenerator().generate_docx_from_html(
        html,
        externalized_table_ids={"large-result-table"},
        externalized_table_titles={"large-result-table": "Large Result Table"},
    )
    with zipfile.ZipFile(io.BytesIO(blob)) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")

    assert "Large table moved to companion Excel appendix" in document_xml
    assert "Large Result Table" in document_xml
    assert "Sample 59" not in document_xml


def test_k_report_does_not_duplicate_method_table():
    generator = ReportGenerator()
    html = generator.generate_k_value_report(
        build_dataset("Sample A"),
        build_results(),
        temperature=20.0,
        porosity=0.35,
        sections={
            "cover_page": False,
            "executive_summary": False,
            "methodology": False,
            "results": True,
            "plots": False,
            "interpretation": False,
            "k_statistics": True,
        },
    )
    table_ids = [table.table_id for table in extract_report_tables(html)]

    assert table_ids.count("k-method-results") == 1
    assert "k-method-details" not in table_ids
